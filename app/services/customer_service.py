from decimal import Decimal

from sqlalchemy import or_, select , func
from sqlalchemy.orm import Session ,selectinload
from sqlalchemy.exc import IntegrityError
from app.models.customer import Customer
from app.models.transaction import Transaction
from app.schemas.customer import CustomerCreate, CustomerUpdate , CustomerSummaryOut

from io import BytesIO , StringIO
from zipfile import BadZipFile
from datetime import date, datetime
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
import csv

def get_latest_customers(db: Session, limit: int = 5):
    stmt = (
        select(Customer)
        .order_by(Customer.id.desc())
        .limit(limit)
    )

    return list(
        db.execute(stmt)
        .scalars()
        .all()
    )


def search_customers(db: Session, keyword: str,limit: int = 20,):
    keyword = keyword.strip()

    if not keyword:
        return []

    conditions = []

    if keyword.isdigit():
        # 數字永遠先比對唯一的客戶 ID
        conditions.append(
            Customer.id == int(keyword)
        )

        # 至少三碼，才搜尋手機尾碼
        if len(keyword) >= 3:
            conditions.append(
                Customer.phone_number.endswith(keyword)
            )
    else:
        # 姓名可輸入一個字或部分姓名
        conditions.append(
            Customer.name.icontains(
                keyword,
                autoescape=True,
            )
        )

    stmt = (
        select(Customer)
        .where(or_(*conditions))
        .order_by(Customer.id.desc())
        .limit(limit)
    )

    return list(
        db.execute(stmt)
        .scalars()
        .all()
    )



def get_customer_by_id(db: Session, customer_id: int):
    stmt = select(Customer).where(
        Customer.id == customer_id
    )

    return (
        db.execute(stmt)
        .scalar_one_or_none()
    )


def create_customer(db: Session, data: CustomerCreate, owner_user_id: int):
    customer = Customer(
        name=data.name,
        phone_number=data.phone_number,
        owner_user_id=owner_user_id,
        gender=data.gender,
        birthday=data.birthday,
        note=data.note,
    )

    db.add(customer)
    db.commit()
    db.refresh(customer)

    return customer


def update_customer(db: Session, customer_id: int, data: CustomerUpdate):
    customer = get_customer_by_id(db, customer_id)

    if customer is None:
        return None

    update_data = data.model_dump(exclude_unset=True)

    for key, value in update_data.items():
        setattr(customer, key, value)

    db.commit()
    db.refresh(customer)

    return customer


def delete_customer(db: Session, customer_id: int):
    customer = get_customer_by_id(db, customer_id)

    if customer is None:
        return None

    db.delete(customer)
    db.commit()

    return customer

def get_customer_summary(
    db: Session,
    customer: Customer,
    limit: int = 5,
) -> CustomerSummaryOut:
    transaction_count, total_amount = db.execute(
        select(
            func.count(Transaction.id),
            func.coalesce(
                func.sum(Transaction.total_amount),
                0,
            ),
        ).where(
            Transaction.customer_id == customer.id
        )
    ).one()

    transaction_count = int(
        transaction_count or 0
    )

    transactions = list(
        db.execute(
            select(Transaction)
            .options(
                selectinload(Transaction.items)
            )
            .where(
                Transaction.customer_id == customer.id
            )
            .order_by(
                Transaction.record_date.desc(),
                Transaction.id.desc(),
            )
            .limit(limit)
        )
        .scalars()
        .all()
    )

    recent_transactions = []

    for index, transaction in enumerate(
        transactions
    ):
        visit_number = transaction_count - index

        items = []

        for item in transaction.items:
            qty = int(item.qty or 0)
            unit_price = Decimal(
                item.unit_price or 0
            )

            subtotal = Decimal(
                item.subtotal
                if item.subtotal is not None
                else qty * unit_price
            )

            items.append({
                "id": item.id,
                "item_name": item.item_name,
                "qty": qty,
                "unit_price": unit_price,
                "subtotal": subtotal,
            })

        recent_transactions.append({
            "id": transaction.id,
            "visit_number": visit_number,
            "total_amount": Decimal(
                transaction.total_amount or 0
            ),
            "note": transaction.note,
            "record_date": transaction.record_date,
            "items": items,
        })

    last_transaction = (
        transactions[0]
        if transactions
        else None
    )

    last_items = None

    if last_transaction:
        item_names = [
            item.item_name
            for item in last_transaction.items
        ]

        if item_names:
            last_items = ", ".join(item_names)

    return CustomerSummaryOut(
        customer_id=customer.id,
        name=customer.name,
        phone_number=customer.phone_number,
        transaction_count=transaction_count,
        total_amount=Decimal(total_amount or 0),
        last_record=(
            last_transaction.note
            if last_transaction
            else None
        ),
        last_items=last_items,
        last_day=(
            last_transaction.record_date
            if last_transaction
            else None
        ),
        recent_transactions=recent_transactions,
    )



def import_customers(
        
    db: Session,
    file_content: bytes,
    owner_user_id: int,
    filename : str
):
    workbook = None
    valid_rows = []
    results =[]
    seen_phones = set()
    
        # 將儲存格內容轉成文字，並去除頭尾空白。
    def cell_text(value):
        return "" if value is None else str(value).strip()
    def add_result(row_number, name, status, reason):
        results.append({
        "row": row_number,
        "name": name,
        "status": status,
        "reason": reason,
    })
    # ===== 1. 依副檔名選擇讀取方式 =====
    extension = filename.lower()

    if extension.endswith(".csv"):
        try:
            # utf-8-sig 同時支援一般 UTF-8 與含 BOM 的 UTF-8。
            text = file_content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV 請使用 UTF-8 編碼") from exc

        rows = csv.reader(
            StringIO(text, newline=""),
            strict=True,
        )

    elif extension.endswith(".xlsx"):
        try:
            workbook = load_workbook(
                BytesIO(file_content),
                read_only=True,
                data_only=True,
            )
        except (
            BadZipFile,
            InvalidFileException,
            OSError,
            ValueError,
            KeyError,
        ) as exc:
            raise ValueError(
                "無法讀取 Excel，請確認檔案是真正的 .xlsx"
            ) from exc

    else:
        raise ValueError("只支援 .csv 或 .xlsx 檔案")

    # ===== 2. 讀取表頭、清理並驗證資料 =====
    # 這個階段不寫入資料庫。
    try:
        if workbook is not None:
            if not workbook.worksheets:
                raise ValueError("Excel 沒有可讀取的工作表")

            # Excel 只讀取第一張工作表。
            rows = workbook.worksheets[0].iter_rows(
                values_only=True
            )

        first_row = next(rows, None)

        if first_row is None:
            raise ValueError("檔案沒有資料")

        headers = [cell_text(value) for value in first_row]

        missing = {"姓名", "電話", "生日"} - set(headers)
        if missing:
            raise ValueError(
                "缺少必要欄位：" + "、".join(sorted(missing))
            )

        nonempty_headers = [header for header in headers if header]
        if len(nonempty_headers) != len(set(nonempty_headers)):
            raise ValueError("表頭不可重複")

        # 第一列是表頭，所以資料從第 2 列開始。
        for row_number, values in enumerate(rows, start=2):
            if row_number > 5001:
                raise ValueError("每次最多匯入 5000 列資料")

            # 完全空白的列直接略過。
            if all(cell_text(value) == "" for value in values):
                continue

            row = dict(zip(headers, values))
            name = cell_text(row.get("姓名"))

            # CSV 維持字串，不會把電話開頭的 0 移除。
            phone = cell_text(row.get("電話"))
            phone = "".join(phone.replace("-", "").split())

            try:
                # 欄數不一致時列為失敗，避免資料錯位或遺失。
                if len(values) != len(headers):
                    raise ValueError("本列欄位數量與表頭不一致")

                birthday = row.get("生日")

                if isinstance(birthday, datetime):
                    birthday = birthday.date()
                elif isinstance(birthday, date):
                    pass
                elif isinstance(birthday, str):
                    #new  # 支援 / 或 - 分隔，例如：
                    # 西元：1995-02-05、1995/2/5
                    # 民國：84/2/5、084-02-05
                    parts = (
                        birthday.strip()
                        .replace("/", "-")
                        .split("-")
                    )

                    if (
                        len(parts) != 3
                        or not all(
                            part.isascii() and part.isdigit()
                            for part in parts
                        )
                        or not 1 <= len(parts[0]) <= 4
                        or not 1 <= len(parts[1]) <= 2
                        or not 1 <= len(parts[2]) <= 2
                    ):
                        raise ValueError(
                            "生日格式錯誤，請輸入 "
                            "1995-02-05 或民國 84/2/5"
                        )

                    year, month, day = map(int, parts)

                    # 規則：1～3 位數年份視為民國；
                    # 4 位數年份視為西元。
                    if len(parts[0]) <= 3:
                        if year < 1:
                            raise ValueError("民國年份必須大於 0")

                        year += 1911

                    try:
                        # date() 也會檢查月份、天數與閏年。
                        birthday = date(year, month, day)
                    except ValueError as exc:
                        raise ValueError("生日不是有效日期") from exc
                    
                # 沿用單筆新增客戶的驗證規則。
                data = CustomerCreate(
                    name=name,
                    phone_number=phone,
                    birthday=birthday,
                    gender=cell_text(row.get("性別")) or None,
                    note=cell_text(row.get("備註")) or None,
                )

                if len(data.name) > 100:
                    raise ValueError("姓名不可超過 100 字")

                if data.gender and len(data.gender) > 10:
                    raise ValueError("性別不可超過 10 字")

            except ValidationError as exc:
                reason = "；".join(
                    f"{'.'.join(map(str, error['loc']))}: "
                    f"{error['msg']}"
                    for error in exc.errors()
                )
                add_result(row_number, name, "failed", reason)
                continue

            except ValueError as exc:
                add_result(row_number, name, "failed", str(exc))
                continue

            # 同一份檔案中，保留第一筆驗證通過的相同電話。
            if data.phone_number in seen_phones:
                add_result(
                    row_number,
                    data.name,
                    "skipped",
                    "同一份檔案內電話重複",
                )
                continue

            seen_phones.add(data.phone_number)
            valid_rows.append((row_number, data))

    except csv.Error as exc:
        raise ValueError(
            "CSV 格式錯誤，請檢查逗號與引號"
        ) from exc

    finally:
        # CSV 沒有 workbook，不能直接呼叫 workbook.close()。
        if workbook is not None:
            workbook.close()

    # ===== 3. 將驗證通過的客戶寫入資料庫 =====
    try:
        for row_number, data in valid_rows:
            exists = db.execute(
                select(Customer.id).where(
                    Customer.phone_number == data.phone_number
                )
            ).scalar_one_or_none()

            if exists is not None:
                add_result(
                    row_number,
                    data.name,
                    "skipped",
                    "資料庫已有相同電話",
                )
                continue

            try:
                # 每筆使用 savepoint：
                # 單筆電話衝突不會撤銷其他已處理的資料。
                with db.begin_nested():
                    customer = Customer(
                        **data.model_dump(),
                        owner_user_id=owner_user_id,
                    )
                    db.add(customer)
                    db.flush()

            except IntegrityError:
                # 防止其他人同時新增相同電話。
                duplicate_id = db.execute(
                    select(Customer.id).where(
                        Customer.phone_number == data.phone_number
                    )
                ).scalar_one_or_none()

                # 不是電話重複造成的資料庫錯誤，就交給外層處理。
                if duplicate_id is None:
                    raise

                add_result(
                    row_number,
                    data.name,
                    "skipped",
                    "資料庫已有相同電話",
                )
                continue

            add_result(
                row_number,
                data.name,
                "created",
                "新增完成",
            )

        db.commit()

    except Exception:
        db.rollback()
        raise

    # ===== 4. 回傳統計與各列結果，供前端顯示 =====
    results.sort(key=lambda item: item["row"])

    return {
        "created": sum(
            item["status"] == "created" for item in results
        ),
        "skipped": sum(
            item["status"] == "skipped" for item in results
        ),
        "failed": sum(
            item["status"] == "failed" for item in results
        ),
        "results": results,
    }